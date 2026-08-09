"""Reproducible micro-benchmarks for local PDF Toolkit core workflows.

The benchmark creates synthetic documents, keeps fixture creation outside the
timed region, validates every output, and reports per-run wall-clock samples.
It intentionally avoids OCR, network access, and user documents.
"""

from __future__ import annotations

import argparse
import cProfile
import json
import os
import platform
import statistics
import sys
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import fitz
from docx import Document
from PIL import Image, ImageDraw

from src.core.batch import HeadlessBatchRunner
from src.handlers.pdf import PDFCompressor
from src.tools.pdf2word.tool import PDFToWordTool


TEXT_PAGES = 120
IMAGE_PAGES = 8
IMAGE_SIZE = (1600, 2200)


@dataclass(frozen=True)
class Fixtures:
    text_pdf: Path
    image_pdf: Path


@dataclass(frozen=True)
class Sample:
    seconds: float
    output_bytes: int


def create_fixtures(root: Path) -> Fixtures:
    """Create deterministic text-heavy and image-heavy PDFs."""
    text_pdf = root / "text-heavy.pdf"
    image_pdf = root / "image-heavy.pdf"

    with fitz.open() as document:
        for page_number in range(TEXT_PAGES):
            page = document.new_page()
            lines = [
                f"Synthetic benchmark page {page_number + 1}, line {line_number + 1}: "
                "local PDF processing performance sample."
                for line_number in range(24)
            ]
            page.insert_textbox(
                fitz.Rect(54, 54, 540, 790),
                "\n".join(lines),
                fontsize=9,
                lineheight=1.25,
            )
        document.save(text_pdf)

    with fitz.open() as document:
        for page_number in range(IMAGE_PAGES):
            image = _synthetic_image(page_number)
            image_path = root / f"fixture-{page_number:02d}.png"
            image.save(image_path, format="PNG", compress_level=6)
            page = document.new_page()
            page.insert_image(fitz.Rect(36, 36, 559, 806), filename=str(image_path))
        document.save(image_pdf)

    return Fixtures(text_pdf=text_pdf, image_pdf=image_pdf)


def _synthetic_image(seed: int) -> Image.Image:
    width, height = IMAGE_SIZE
    image = Image.new("RGB", IMAGE_SIZE, (245, 245, 240))
    draw = ImageDraw.Draw(image)
    for y in range(0, height, 24):
        color = (
            (seed * 31 + y // 3) % 256,
            (seed * 53 + y // 5 + 80) % 256,
            (seed * 71 + y // 7 + 160) % 256,
        )
        draw.rectangle((0, y, width, min(y + 23, height)), fill=color)
    for x in range(0, width, 80):
        draw.line((x, 0, width - x // 2, height), fill=(255, 255, 255), width=3)
    for index in range(40):
        x = (index * 97 + seed * 43) % (width - 260)
        y = (index * 151 + seed * 67) % (height - 120)
        draw.rectangle((x, y, x + 240, y + 80), outline=(20, 20, 20), width=4)
        draw.text((x + 10, y + 10), f"page={seed + 1} block={index + 1}", fill=(10, 10, 10))
    return image


def benchmark_pdf_to_images(fixtures: Fixtures, output_dir: Path) -> int:
    outputs = HeadlessBatchRunner.convert_pdf_to_images(
        str(fixtures.image_pdf), str(output_dir), 150, "png"
    )
    if len(outputs) != IMAGE_PAGES:
        raise RuntimeError(f"Expected {IMAGE_PAGES} rendered pages, got {len(outputs)}.")
    return sum(Path(path).stat().st_size for path in outputs)


def benchmark_pdf_to_word_text(fixtures: Fixtures, output_dir: Path) -> int:
    output_path = output_dir / "text-heavy.docx"
    PDFToWordTool.convert_pdf_to_docx(
        str(fixtures.text_pdf), str(output_path), mode="Text Only"
    )
    document = Document(output_path)
    if len(document.paragraphs) <= TEXT_PAGES:
        raise RuntimeError("Text-only DOCX did not contain the expected page content.")
    return output_path.stat().st_size


def benchmark_pdf_to_word_page_images(fixtures: Fixtures, output_dir: Path) -> int:
    output_path = output_dir / "image-heavy.docx"
    PDFToWordTool.convert_pdf_to_docx(
        str(fixtures.image_pdf), str(output_path), mode="Page Images"
    )
    document = Document(output_path)
    if len(document.inline_shapes) != IMAGE_PAGES:
        raise RuntimeError(
            f"Expected {IMAGE_PAGES} DOCX page images, got {len(document.inline_shapes)}."
        )
    return output_path.stat().st_size


def benchmark_compress_pdf(fixtures: Fixtures, output_dir: Path) -> int:
    output_path = output_dir / "image-heavy-compressed.pdf"
    compressor = PDFCompressor(level="medium")
    if not compressor.compress(str(fixtures.image_pdf), str(output_path)):
        raise RuntimeError("Synthetic PDF compression failed.")
    with fitz.open(output_path) as document:
        if document.page_count != IMAGE_PAGES:
            raise RuntimeError(
                f"Expected {IMAGE_PAGES} compressed pages, got {document.page_count}."
            )
    return output_path.stat().st_size


SCENARIOS: dict[str, Callable[[Fixtures, Path], int]] = {
    "pdf_to_images": benchmark_pdf_to_images,
    "pdf_to_word_text": benchmark_pdf_to_word_text,
    "pdf_to_word_page_images": benchmark_pdf_to_word_page_images,
    "compress_pdf": benchmark_compress_pdf,
}


def run_sample(
    scenario: Callable[[Fixtures, Path], int], fixtures: Fixtures, root: Path, name: str
) -> Sample:
    output_dir = root / name
    output_dir.mkdir(parents=True)
    started = time.perf_counter()
    output_bytes = scenario(fixtures, output_dir)
    seconds = time.perf_counter() - started
    return Sample(seconds=seconds, output_bytes=output_bytes)


def summarize(samples: list[Sample]) -> dict[str, object]:
    timings = [sample.seconds for sample in samples]
    output_sizes = {sample.output_bytes for sample in samples}
    if len(output_sizes) != 1:
        raise RuntimeError(f"Output size changed between samples: {sorted(output_sizes)}")
    return {
        "median_seconds": statistics.median(timings),
        "min_seconds": min(timings),
        "max_seconds": max(timings),
        "samples_seconds": timings,
        "output_bytes": output_sizes.pop(),
    }


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--scenario",
        action="append",
        choices=tuple(SCENARIOS),
        help="Scenario to run; repeat the option to select multiple scenarios.",
    )
    parser.add_argument("--repeat", type=int, default=3, help="Measured runs per scenario.")
    parser.add_argument("--warmup", type=int, default=1, help="Untimed warm-up runs per scenario.")
    parser.add_argument("--profile", choices=tuple(SCENARIOS), help="Also profile one run.")
    parser.add_argument("--profile-out", type=Path, help="Path for cProfile statistics.")
    parser.add_argument("--json-out", type=Path, help="Also write the JSON report to this path.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.repeat < 1 or args.warmup < 0:
        raise SystemExit("--repeat must be at least 1 and --warmup cannot be negative.")
    if args.profile_out and not args.profile:
        raise SystemExit("--profile-out requires --profile.")

    selected = args.scenario or list(SCENARIOS)
    with tempfile.TemporaryDirectory(prefix="pdf-toolkit-benchmark-") as temporary:
        root = Path(temporary)
        fixtures = create_fixtures(root)
        results: dict[str, object] = {}

        for scenario_name in selected:
            scenario = SCENARIOS[scenario_name]
            for warmup_index in range(args.warmup):
                run_sample(scenario, fixtures, root, f"{scenario_name}-warmup-{warmup_index}")
            samples = [
                run_sample(scenario, fixtures, root, f"{scenario_name}-run-{run_index}")
                for run_index in range(args.repeat)
            ]
            results[scenario_name] = summarize(samples)

        if args.profile:
            profile_path = args.profile_out or Path(f"{args.profile}.prof")
            profile_path.parent.mkdir(parents=True, exist_ok=True)
            profiler = cProfile.Profile()
            profiler.enable()
            run_sample(SCENARIOS[args.profile], fixtures, root, f"{args.profile}-profile")
            profiler.disable()
            profiler.dump_stats(profile_path)

        payload = {
            "schema_version": 1,
            "environment": {
                "python": sys.version.split()[0],
                "platform": platform.platform(),
                "processor": platform.processor(),
                "cpu_count": os.cpu_count(),
                "pymupdf": fitz.VersionBind,
                "pillow": Image.__version__,
            },
            "fixture": {
                "text_pages": TEXT_PAGES,
                "image_pages": IMAGE_PAGES,
                "image_size": list(IMAGE_SIZE),
            },
            "repeat": args.repeat,
            "warmup": args.warmup,
            "results": results,
        }
        rendered = json.dumps(payload, indent=2, sort_keys=True)
        print(rendered)
        if args.json_out:
            args.json_out.parent.mkdir(parents=True, exist_ok=True)
            args.json_out.write_text(rendered + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
